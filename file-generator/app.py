import os
import uuid
from flask import Flask, request, jsonify, send_from_directory
import requests
from fpdf import FPDF
from docx import Document
from openpyxl import Workbook
from flasgger import Swagger, swag_from

app = Flask(__name__)

# Swagger configuration
app.config['SWAGGER'] = {
    'title': 'Dify File Generator API',
    'uiversion': 3,
    'specs_route': '/apidocs/',
    'description': 'API for generating files (PDF, DOCX, XLSX) and connecting to Dify agent',
    'version': '1.0.2'
}
swagger = Swagger(app)

# Configuration
GENERATED_FILES_DIR = os.path.join('static', 'generated_files')
if not os.path.exists(GENERATED_FILES_DIR):
    os.makedirs(GENERATED_FILES_DIR)

# --- Helper Functions for File Generation ---

def create_pdf(content, filename):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    # Add content to the PDF, handling multi-line text
    for line in content.split('\n'):
        pdf.multi_cell(0, 10, txt=line)
    
    filepath = os.path.join(GENERATED_FILES_DIR, filename)
    pdf.output(filepath)
    return filepath

def create_docx(content, filename):
    document = Document()
    document.add_paragraph(content)
    filepath = os.path.join(GENERATED_FILES_DIR, filename)
    document.save(filepath)
    return filepath

def create_xlsx(content, filename):
    workbook = Workbook()
    sheet = workbook.active
    # Split content by lines and then by a delimiter (e.g., comma) for cells
    for r_idx, row_data in enumerate(content.split('\n'), 1):
        for c_idx, cell_value in enumerate(row_data.split(','), 1):
            sheet.cell(row=r_idx, column=c_idx, value=cell_value.strip())

    filepath = os.path.join(GENERATED_FILES_DIR, filename)
    workbook.save(filepath)
    return filepath

# --- API Endpoints ---

@app.route('/anythingllm-manifest.json', methods=['GET'])
@swag_from({
    "responses": {
        200: {
            "description": "AnythingLLM manifest file",
            "schema": {
                "type": "object",
                "properties": {
                    "version": {"type": "string"},
                    "name": {"type": "string"},
                    "description": {"type": "string"},
                    "type": {"type": "string"},
                    "entrypoint": {"type": "string"},
                    "schemas": {"type": "array"}
                }
            }
        }
    }
})
def manifest():
    return jsonify({
        "version": "1.0.2",
        "name": "Dify File Generator",
        "description": "An agent that can generate files (PDF, DOCX, XLSX), save chat history, or save the last response using a Dify agent.",
        "type": "agent",
        "entrypoint": f"{request.host_url}invoke",
        "schemas": [
            {
                "type": "agent",
                "name": "dify-file-generator",
                "provider": "dify",
                "description": "Use this agent to create files from content, save the full conversation, or save the last response.",
                "settings": [
                    {
                        "name": "Dify API Endpoint",
                        "type": "text",
                        "description": "Your Dify application's API endpoint (e.g., https://api.dify.ai/v1)",
                        "value": "",
                        "required": True
                    },
                    {
                        "name": "Dify API Secret Key",
                        "type": "secret",
                        "description": "Your Dify application's secret key.",
                        "value": "",
                        "required": True
                    }
                ],
                "examples": [
                    "Create a PDF named 'report.pdf' with the content 'This is a test'.",
                    "Save this conversation as a DOCX file named 'chat_log.docx'.",
                    "Save the last response as a file named 'summary.pdf'."
                ]
            }
        ]
    })

@app.route('/generate/<file_type>', methods=['POST'])
@swag_from({
    "parameters": [
        {
            "name": "file_type",
            "in": "path",
            "type": "string",
            "enum": ["pdf", "docx", "xlsx"],
            "required": True,
            "description": "Type of file to generate"
        },
        {
            "name": "body",
            "in": "body",
            "schema": {
                "type": "object",
                "properties": {
                    "content": {"type": "string", "description": "Content to include in the file"},
                    "filename": {"type": "string", "description": "Name of the file to generate"}
                },
                "required": ["content", "filename"]
            }
        }
    ],
    "responses": {
        200: {
            "description": "File generated successfully",
            "schema": {
                "type": "object",
                "properties": {
                    "url": {"type": "string", "description": "URL to access the generated file"}
                }
            }
        },
        400: {
            "description": "Bad request",
            "schema": {
                "type": "object",
                "properties": {
                    "error": {"type": "string"}
                }
            }
        },
        500: {
            "description": "Internal server error",
            "schema": {
                "type": "object",
                "properties": {
                    "error": {"type": "string"}
                }
            }
        }
    }
})
def generate_file(file_type):
    data = request.get_json()
    content = data.get('content', '')
    filename = data.get('filename')

    if not filename:
        return jsonify({"error": "Filename is required"}), 400

    try:
        if file_type == 'pdf':
            filepath = create_pdf(content, filename)
        elif file_type == 'docx':
            filepath = create_docx(content, filename)
        elif file_type == 'xlsx':
            filepath = create_xlsx(content, filename)
        else:
            return jsonify({"error": "Unsupported file type"}), 400
        
        file_url = f"{request.host_url}{filepath}"
        return jsonify({"url": file_url})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/invoke', methods=['POST'])
@swag_from({
    "parameters": [
        {
            "name": "body",
            "in": "body",
            "schema": {
                "type": "object",
                "properties": {
                    "prompt": {"type": "string", "description": "User's prompt to send to Dify"},
                    "history": {
                        "type": "array",
                        "description": "Conversation history",
                        "items": {
                            "type": "object",
                            "properties": {
                                "role": {"type": "string", "enum": ["user", "assistant"]},
                                "content": {"type": "string"}
                            }
                        }
                    },
                    "settings": {
                        "type": "object",
                        "properties": {
                            "Dify API Endpoint": {"type": "string"},
                            "Dify API Secret Key": {"type": "string"}
                        }
                    },
                    "conversation_id": {"type": "string", "description": "Optional conversation ID"}
                },
                "required": ["prompt", "settings"]
            }
        }
    ],
    "responses": {
        200: {
            "description": "Response from Dify agent",
            "schema": {
                "type": "object",
                "properties": {
                    "text": {"type": "string", "description": "Response text from Dify agent"}
                }
            }
        }
    }
})
def invoke_dify_agent():
    data = request.get_json()
    user_prompt = data.get('prompt', '')
    conversation_history = data.get('history', [])
    settings = data.get('settings', {})

    dify_api_endpoint = settings.get('Dify API Endpoint')
    dify_api_key = settings.get('Dify API Secret Key')

    if not dify_api_endpoint or not dify_api_key:
        return jsonify({"text": "Dify API credentials are not configured."})

    # --- Process History for Dify --- #
    # Find the last assistant message
    last_assistant_message = ""
    for message in reversed(conversation_history):
        if message.get('role') == 'assistant':
            last_assistant_message = message.get('content', '')
            break

    # Format the full conversation history into a single string
    formatted_chat_history = ""
    for message in conversation_history:
        role = "User" if message.get('role') == 'user' else "Assistant"
        content = message.get('content', '')
        formatted_chat_history += f"{role}: {content}\n\n"
    # --- End History Processing --- #

    headers = {
        'Authorization': f'Bearer {dify_api_key}',
        'Content-Type': 'application/json'
    }

    payload = {
        'inputs': {
            'last_assistant_message': last_assistant_message,
            'formatted_chat_history': formatted_chat_history.strip()
        },
        'query': user_prompt,
        'response_mode': 'streaming',
        'user': f'anythingllm-user-{uuid.uuid4()}',
        'conversation_id': data.get('conversation_id', ''),
        'files': []
    }

    try:
        response = requests.post(
            f"{dify_api_endpoint}/chat-messages",
            headers=headers,
            json=payload,
            stream=True
        )
        response.raise_for_status()

        # Process streaming response
        final_answer = ""
        for line in response.iter_lines():
            if line:
                decoded_line = line.decode('utf-8')
                if decoded_line.startswith('data:'):
                    import json
                    try:
                        json_data = json.loads(decoded_line[len('data: '):])
                        if json_data.get('event') == 'agent_message':
                            final_answer += json_data.get('answer', '')
                    except json.JSONDecodeError:
                        continue

        return jsonify({"text": final_answer})

    except requests.exceptions.RequestException as e:
        return jsonify({"text": f"Error calling Dify agent: {str(e)}"})

@app.route('/static/generated_files/<filename>')
@swag_from({
    "parameters": [
        {
            "name": "filename",
            "in": "path",
            "type": "string",
            "required": True,
            "description": "Name of the file to serve"
        }
    ],
    "responses": {
        200: {
            "description": "File content"
        },
        404: {
            "description": "File not found"
        }
    }
})
def serve_generated_file(filename):
    return send_from_directory(GENERATED_FILES_DIR, filename)

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5001, debug=True)
